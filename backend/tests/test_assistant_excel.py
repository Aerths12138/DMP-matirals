"""Focused checks for bounded Excel conversations and import preparation."""

import json
from io import BytesIO

import pytest
from fastapi import UploadFile
from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import selectinload

from app.models.material import Material
from app.models.permission import Permission, Role
from app.models.property import Property
from app.models.test import Test as MaterialTest
from app.models.test import TestMeasurement as MaterialTestMeasurement
from app.models.user import User
from app.services import assistant_excel_service, assistant_service


EXCEL_CHAT_URL = "/api/v1/assistant/chat/excel"
HISTORY_URL = "/api/v1/assistant/history"


def _workbook_bytes(*, sentinel: str = "RAW-CELL-SENTINEL") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "材料数据"
    sheet.append(["材料名称", "电导率", "备注"])
    sheet.append(["演示合金", "1.20E-5", sentinel])
    second = workbook.create_sheet("测试条件")
    second.append(["温度", "单位"])
    second.append([25, "°C"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _multipart_request(content: str) -> dict[str, str]:
    return {
        "request": json.dumps(
            {
                "messages": [{"role": "user", "content": content}],
                "context_path": "/assistant",
            },
            ensure_ascii=False,
        )
    }


async def _grant_import_permission(api, user_id) -> None:
    async with api.session_factory() as session:
        user = await session.scalar(
            select(User)
            .options(selectinload(User.role).selectinload(Role.permissions))
            .where(User.id == user_id)
        )
        permission = await session.scalar(
            select(Permission).where(Permission.codename == "import:create")
        )
        if permission is None:
            permission = Permission(codename="import:create")
            session.add(permission)
            await session.flush()
        if permission not in user.role.permissions:
            user.role.permissions.append(permission)
        await session.commit()


def _import_workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "介电性能"
    sheet.append(["材料名称", "材料类型", "电导率", "温度", "批次说明"])
    sheet.append(["格式化合金", "金属", "1.20E-5", "25", "首批"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _two_sheet_import_workbook_bytes() -> bytes:
    workbook = Workbook()
    first = workbook.active
    first.title = "首表数据"
    first.append(["材料名称", "材料类型", "电导率"])
    first.append(["首表材料", "金属", "1E-5"])
    second = workbook.create_sheet("目标表数据")
    second.append(["材料名称", "材料类型", "电导率"])
    second.append(["第二表材料", "陶瓷", "2E-5"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


@pytest.mark.asyncio
async def test_excel_chat_passes_bounded_preview_without_persisting_cells(
    api,
    monkeypatch,
):
    sentinel = "RAW-CELL-NOT-IN-HISTORY"
    captured: dict = {}

    async def fake_llm(messages, tools):
        captured["messages"] = messages
        captured["tools"] = tools
        return {"role": "assistant", "content": "已完成只读分析。"}

    monkeypatch.setattr(assistant_service, "request_llm", fake_llm)
    response = await api.client.post(
        EXCEL_CHAT_URL,
        data=_multipart_request("请概述这个工作簿。[Excel：demo.xlsx]"),
        files={
            "file": (
                "demo.xlsx",
                _workbook_bytes(sentinel=sentinel),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=api.headers["researcher"],
    )

    assert response.status_code == 200, response.text
    assert response.json()["message"] == "已完成只读分析。"
    assert response.json()["pending_actions"] == []

    attachment_context = next(
        item["content"]
        for item in captured["messages"]
        if item["role"] == "system" and sentinel in item["content"]
    )
    assert "不可信数据" in attachment_context
    assert "1.20E-5" in attachment_context
    assert '"sheet_count": 2' in attachment_context
    advertised_tools = {
        item["function"]["name"] for item in (captured["tools"] or [])
    }
    assert advertised_tools.isdisjoint(assistant_service.WRITE_TOOL_NAMES)

    history = await api.client.get(HISTORY_URL, headers=api.headers["researcher"])
    assert history.status_code == 200
    serialized_history = json.dumps(history.json(), ensure_ascii=False)
    assert sentinel not in serialized_history
    assert "1.20E-5" not in serialized_history
    assert history.json()["messages"] == [
        {"role": "user", "content": "请概述这个工作簿。[Excel：demo.xlsx]"},
        {"role": "assistant", "content": "已完成只读分析。"},
    ]


@pytest.mark.asyncio
async def test_excel_import_preparation_is_deterministic_and_requires_confirmation(
    api,
    monkeypatch,
):
    await _grant_import_permission(api, api.user_ids["researcher"])
    calls = 0

    async def fake_llm(messages, tools):
        nonlocal calls
        calls += 1
        if calls == 1:
            advertised_tools = {
                item["function"]["name"] for item in (tools or [])
            }
            assert "prepare_excel_import" in advertised_tools
            return {
                "role": "assistant",
                "content": None,
                "tool_calls": [
                    {
                        "id": "prepare-import",
                        "type": "function",
                        "function": {
                            "name": "prepare_excel_import",
                            "arguments": "{}",
                        },
                    }
                ],
            }
        assert messages[-1]["role"] == "tool"
        tool_result = json.loads(messages[-1]["content"])
        assert tool_result["ok"] is True
        return {"role": "assistant", "content": "格式化预览已生成，请确认后导入。"}

    monkeypatch.setattr(assistant_service, "request_llm", fake_llm)
    response = await api.client.post(
        EXCEL_CHAT_URL,
        data=_multipart_request("请按平台格式整理这个文件并准备导入。"),
        files={
            "file": (
                "介电性能.xlsx",
                _import_workbook_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=api.headers["researcher"],
    )

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["message"] == "格式化预览已生成，请确认后导入。"
    assert body["pending_actions"] == []
    preview = body["pending_import"]
    assert preview["mapping"] == {
        "材料名称": "name",
        "材料类型": "material_type",
        "电导率": "result",
        "温度": "condition",
        "批次说明": "property",
    }
    assert preview["normalized_rows"] == [
        ["格式化合金", "metal", 0.000012, "25", "首批"]
    ]
    assert preview["conversion_count"] == 2
    assert len(preview["import_rows"]) == 1

    import_row = preview["import_rows"][0]
    assert import_row["name"] == "格式化合金"
    assert import_row["material_type"] == "metal"
    assert import_row["properties"] == {"批次说明": "首批"}
    assert len(import_row["tests"]) == 1
    test_row = import_row["tests"][0]
    assert test_row["test_name"] == "导入测试"
    assert test_row["conditions"] == {"温度": "25"}
    assert len(test_row["measurements"]) == 1
    measurement_row = test_row["measurements"][0]
    assert measurement_row["property_name"] == "电导率"
    assert measurement_row["value_number"] == pytest.approx(0.000012)

    async with api.session_factory() as session:
        before_confirmation = await session.scalar(
            select(Material).where(Material.name == "格式化合金")
        )
    assert before_confirmation is None

    confirmed = await api.client.post(
        "/api/v1/materials/import",
        headers=api.headers["researcher"],
        json={"rows": preview["import_rows"]},
    )
    assert confirmed.status_code == 201, confirmed.text
    assert confirmed.json() == {"total": 1, "created": 1, "failed": 0, "errors": []}

    async with api.session_factory() as session:
        material = await session.scalar(
            select(Material)
            .options(
                selectinload(Material.properties),
                selectinload(Material.tests).selectinload(MaterialTest.measurements),
            )
            .where(Material.name == "格式化合金")
        )
        properties = await session.scalars(
            select(Property).where(Property.material_id == material.id)
        )
        measurements = await session.scalars(
            select(MaterialTestMeasurement)
            .join(MaterialTest)
            .where(MaterialTest.material_id == material.id)
        )

    assert material.material_type == "metal"
    assert [(item.property_name, item.value_text) for item in properties] == [
        ("批次说明", "首批")
    ]
    assert [
        (item.property_name, item.value_number) for item in measurements
    ] == [("电导率", pytest.approx(0.000012))]


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("sheet_name", "expected_sheet", "expected_material", "excluded_material"),
    [
        (None, "首表数据", "首表材料", "第二表材料"),
        ("目标表数据", "目标表数据", "第二表材料", "首表材料"),
    ],
)
async def test_excel_import_prepares_selected_sheet_and_defaults_to_first(
    api,
    monkeypatch,
    sheet_name,
    expected_sheet,
    expected_material,
    excluded_material,
):
    await _grant_import_permission(api, api.user_ids["researcher"])
    calls = 0
    captured_attachment = ""

    async def fake_llm(messages, tools):
        nonlocal calls, captured_attachment
        calls += 1
        if calls == 1:
            captured_attachment = next(
                item["content"]
                for item in messages
                if item["role"] == "system" and '"sheet_count": 2' in item["content"]
            )
            assert "prepare_excel_import" in {
                item["function"]["name"] for item in (tools or [])
            }
            return {
                "role": "assistant",
                "content": None,
                "tool_calls": [
                    {
                        "id": "prepare-selected-sheet",
                        "type": "function",
                        "function": {
                            "name": "prepare_excel_import",
                            "arguments": "{}",
                        },
                    }
                ],
            }
        tool_result = json.loads(messages[-1]["content"])
        assert tool_result["ok"] is True
        assert tool_result["data"]["sheet_name"] == expected_sheet
        return {"role": "assistant", "content": "已生成指定工作表的导入预览。"}

    monkeypatch.setattr(assistant_service, "request_llm", fake_llm)
    request_data = _multipart_request("格式化指定工作表并准备导入。")
    if sheet_name is not None:
        request_data["sheet_name"] = sheet_name
    response = await api.client.post(
        EXCEL_CHAT_URL,
        data=request_data,
        files={"file": ("多表数据.xlsx", _two_sheet_import_workbook_bytes())},
        headers=api.headers["researcher"],
    )

    assert response.status_code == 200, response.text
    preview = response.json()["pending_import"]
    assert preview["sheet_name"] == expected_sheet
    assert [row["name"] for row in preview["import_rows"]] == [expected_material]
    assert excluded_material not in json.dumps(preview, ensure_ascii=False)
    if sheet_name is not None:
        assert captured_attachment.index(expected_sheet) < captured_attachment.index(
            "首表数据"
        )


@pytest.mark.asyncio
async def test_excel_chat_rejects_unknown_sheet_before_model(api, monkeypatch):
    async def unexpected_llm(*_args, **_kwargs):
        raise AssertionError("An unknown worksheet must be rejected before the model")

    monkeypatch.setattr(assistant_service, "request_llm", unexpected_llm)
    response = await api.client.post(
        EXCEL_CHAT_URL,
        data={
            **_multipart_request("读取并准备指定工作表。"),
            "sheet_name": "不存在的工作表",
        },
        files={"file": ("多表数据.xlsx", _two_sheet_import_workbook_bytes())},
        headers=api.headers["researcher"],
    )

    assert response.status_code == 422, response.text
    assert "不存在的工作表" in response.text


@pytest.mark.asyncio
async def test_excel_import_preparation_is_not_available_without_permission(
    api,
    monkeypatch,
):
    calls = 0

    async def fake_llm(messages, tools):
        nonlocal calls
        calls += 1
        if calls == 1:
            advertised_tools = {
                item["function"]["name"] for item in (tools or [])
            }
            assert "prepare_excel_import" not in advertised_tools
            return {
                "role": "assistant",
                "content": None,
                "tool_calls": [
                    {
                        "id": "forged-prepare-import",
                        "type": "function",
                        "function": {
                            "name": "prepare_excel_import",
                            "arguments": "{}",
                        },
                    }
                ],
            }
        tool_result = json.loads(messages[-1]["content"])
        assert tool_result["ok"] is False
        assert tool_result["status"] == 403
        return {"role": "assistant", "content": "你没有数据导入权限。"}

    monkeypatch.setattr(assistant_service, "request_llm", fake_llm)
    response = await api.client.post(
        EXCEL_CHAT_URL,
        data=_multipart_request("整理后导入这个文件。"),
        files={"file": ("no-access.xlsx", _import_workbook_bytes())},
        headers=api.headers["viewer"],
    )

    assert response.status_code == 200, response.text
    assert response.json()["pending_actions"] == []
    assert response.json()["pending_import"] is None
    async with api.session_factory() as session:
        created = await session.scalar(
            select(Material).where(Material.name == "格式化合金")
        )
    assert created is None


@pytest.mark.asyncio
async def test_excel_preview_is_bounded_and_reports_truncation():
    workbook = Workbook()
    first = workbook.active
    first.title = "超限数据"
    for row in range(1, 32):
        for column in range(1, 22):
            first.cell(row=row, column=column, value=f"R{row}C{column}")
    first.cell(row=2, column=2, value="X" * 250)
    for index in range(2, 5):
        workbook.create_sheet(f"工作表{index}").append([index])
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    preview = await assistant_excel_service.parse_excel_upload(
        UploadFile(filename="bounded.xlsx", file=BytesIO(output.getvalue()))
    )

    assert preview["sheet_count"] == 4
    assert preview["previewed_sheet_count"] == 3
    assert preview["truncated"] is True
    first_preview = preview["sheets"][0]
    assert first_preview["total_rows"] == 31
    assert first_preview["total_columns"] == 21
    assert len(first_preview["preview_rows"]) == assistant_excel_service.MAX_ROWS
    assert all(
        len(row) <= assistant_excel_service.MAX_COLUMNS
        for row in first_preview["preview_rows"]
    )
    serialized_preview = json.dumps(preview, ensure_ascii=False)
    assert "R31C1" not in serialized_preview
    assert "R1C21" not in serialized_preview
    assert "X" * 201 not in serialized_preview


@pytest.mark.asyncio
async def test_excel_import_preparation_rejects_more_than_500_data_rows(
    api,
    monkeypatch,
):
    await _grant_import_permission(api, api.user_ids["researcher"])
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "超限导入"
    sheet.append(["材料名称", "电导率"])
    for index in range(501):
        sheet.append([f"材料-{index}", "1E-5"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    async def unexpected_llm(*_args, **_kwargs):
        raise AssertionError("An oversized import sheet must be rejected before the model")

    monkeypatch.setattr(assistant_service, "request_llm", unexpected_llm)
    response = await api.client.post(
        EXCEL_CHAT_URL,
        data=_multipart_request("格式化并准备导入。"),
        files={"file": ("too-many-rows.xlsx", output.getvalue())},
        headers=api.headers["researcher"],
    )

    assert response.status_code == 422, response.text
    assert "500" in response.text
    async with api.session_factory() as session:
        material_count = len((await session.scalars(select(Material))).all())
    assert material_count == 0


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("filename", "content", "expected_status"),
    [
        ("legacy.xls", b"legacy", 415),
        ("damaged.xlsx", b"not-a-zip", 422),
    ],
)
async def test_excel_chat_rejects_unsupported_or_damaged_files(
    api,
    monkeypatch,
    filename,
    content,
    expected_status,
):
    async def unexpected_llm(*_args, **_kwargs):
        raise AssertionError("The model must not receive an invalid workbook")

    monkeypatch.setattr(assistant_service, "request_llm", unexpected_llm)
    response = await api.client.post(
        EXCEL_CHAT_URL,
        data=_multipart_request("读取文件"),
        files={"file": (filename, content, "application/octet-stream")},
        headers=api.headers["researcher"],
    )

    assert response.status_code == expected_status
    assert "Traceback" not in response.text
    history = await api.client.get(HISTORY_URL, headers=api.headers["researcher"])
    assert history.status_code == 200
    assert history.json()["messages"] == []


@pytest.mark.asyncio
async def test_excel_mode_rejects_unadvertised_write_tool_calls(api, monkeypatch):
    calls = 0

    async def fake_llm(messages, tools):
        nonlocal calls
        calls += 1
        if calls == 1:
            assert all(
                item["function"]["name"] != "create_material"
                for item in (tools or [])
            )
            return {
                "role": "assistant",
                "content": None,
                "tool_calls": [
                    {
                        "id": "forged-write",
                        "type": "function",
                        "function": {
                            "name": "create_material",
                            "arguments": json.dumps(
                                {"name": "MALICIOUS-FROM-CELL", "material_type": "metal"}
                            ),
                        },
                    }
                ],
            }
        tool_result = json.loads(messages[-1]["content"])
        assert tool_result["ok"] is False
        assert tool_result["status"] == 403
        assert "read-only" in tool_result["error"]
        return {"role": "assistant", "content": "附件模式只进行读取分析。"}

    monkeypatch.setattr(assistant_service, "request_llm", fake_llm)
    response = await api.client.post(
        EXCEL_CHAT_URL,
        data=_multipart_request("读取这个文件"),
        files={"file": ("safe.xlsx", _workbook_bytes(), "application/octet-stream")},
        headers=api.headers["researcher"],
    )

    assert response.status_code == 200, response.text
    assert response.json()["pending_actions"] == []
    async with api.session_factory() as session:
        created = await session.scalar(
            select(Material).where(Material.name == "MALICIOUS-FROM-CELL")
        )
    assert created is None

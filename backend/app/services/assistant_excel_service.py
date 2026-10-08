"""Bounded, in-memory Excel previews for AI assistant conversations."""

import math
from datetime import date, datetime, time
from io import BytesIO
from pathlib import PurePath
from typing import Any
from zipfile import BadZipFile, ZipFile

from fastapi import HTTPException, UploadFile, status
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from starlette.concurrency import run_in_threadpool


MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_EXPANDED_BYTES = 50 * 1024 * 1024
MAX_ZIP_ENTRIES = 1000
MAX_SHEETS = 3
MAX_ROWS = 30
MAX_COLUMNS = 20
MAX_CELL_CHARACTERS = 200
MAX_TOTAL_TEXT_CHARACTERS = 20_000
MAX_IMPORT_ROWS = 500
MAX_IMPORT_COLUMNS = 100
MAX_IMPORT_CELL_CHARACTERS = 20_000
MAX_IMPORT_TOTAL_TEXT_CHARACTERS = 1_000_000


def _safe_filename(filename: str | None) -> str:
    safe_name = PurePath((filename or "").replace("\\", "/")).name
    return "".join(
        character if character >= " " and character != "\x7f" else "_"
        for character in safe_name
    )[:500] or "unknown.xlsx"


def _validate_xlsx_archive(content: bytes) -> None:
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            expanded_bytes = sum(item.file_size for item in entries)
    except BadZipFile as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="上传的文件不是有效的 .xlsx 工作簿",
        ) from exc

    if len(entries) > MAX_ZIP_ENTRIES or expanded_bytes > MAX_EXPANDED_BYTES:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail="Excel 工作簿解压后超过安全预览限制",
        )


def _cell_value(value: Any, remaining_characters: int) -> tuple[Any, int, bool, bool]:
    """Return a JSON-safe cell value plus budget and truncation metadata."""
    if value is None:
        return None, 0, False, False
    if isinstance(value, bool):
        text_length = len(str(value))
        if text_length > remaining_characters:
            return None, 0, False, True
        return value, text_length, False, False
    if isinstance(value, int):
        text_length = len(str(value))
        if text_length > remaining_characters:
            return None, 0, False, True
        return value, text_length, False, False
    if isinstance(value, float):
        normalized: float | str = value if math.isfinite(value) else str(value)
        text_length = len(str(normalized))
        if text_length > remaining_characters:
            return None, 0, False, True
        return normalized, text_length, False, False
    if isinstance(value, (datetime, date, time)):
        text = value.isoformat()
    else:
        text = str(value)

    cell_truncated = len(text) > MAX_CELL_CHARACTERS
    text = text[:MAX_CELL_CHARACTERS]
    if not text:
        return "", 0, cell_truncated, False
    if remaining_characters <= 0:
        return None, 0, cell_truncated, True
    if len(text) > remaining_characters:
        return text[:remaining_characters], remaining_characters, True, True
    return text, len(text), cell_truncated, False


def _selected_worksheet(workbook, sheet_name: str | None):
    if not workbook.worksheets:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Excel 工作簿没有可用的工作表",
        )
    if sheet_name is None:
        return workbook.worksheets[0]
    for worksheet in workbook.worksheets:
        if worksheet.title == sheet_name:
            return worksheet
    raise HTTPException(
        status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
        detail=f"指定的 Excel 工作表“{sheet_name}”不存在，请重新选择",
    )


def _parse_excel_content(
    content: bytes,
    filename: str,
    sheet_name: str | None = None,
) -> dict:
    _validate_xlsx_archive(content)

    workbook = None
    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        selected_worksheet = _selected_worksheet(workbook, sheet_name)
        worksheets = [
            selected_worksheet,
            *(item for item in workbook.worksheets if item is not selected_worksheet),
        ]
        sheet_count = len(workbook.worksheets)
        remaining_characters = MAX_TOTAL_TEXT_CHARACTERS
        sheets: list[dict] = []
        workbook_truncated = sheet_count > MAX_SHEETS

        for worksheet in worksheets[:MAX_SHEETS]:
            total_rows = int(worksheet.max_row or 0)
            total_columns = int(worksheet.max_column or 0)
            sheet_truncated = total_rows > MAX_ROWS or total_columns > MAX_COLUMNS
            preview_rows: list[list[Any]] = []

            row_limit = min(total_rows, MAX_ROWS)
            column_limit = min(total_columns, MAX_COLUMNS)
            budget_exhausted = False

            if row_limit and column_limit:
                for values in worksheet.iter_rows(
                    min_row=1,
                    max_row=row_limit,
                    min_col=1,
                    max_col=column_limit,
                    values_only=True,
                ):
                    preview_row: list[Any] = []
                    for value in values:
                        normalized, used, cell_truncated, exhausted = _cell_value(
                            value,
                            remaining_characters,
                        )
                        if exhausted and normalized is None:
                            budget_exhausted = True
                            break
                        preview_row.append(normalized)
                        remaining_characters -= used
                        if cell_truncated:
                            sheet_truncated = True
                        if exhausted:
                            budget_exhausted = True
                            break
                    if preview_row:
                        preview_rows.append(preview_row)
                    if budget_exhausted:
                        sheet_truncated = True
                        break

            sheets.append(
                {
                    "name": worksheet.title,
                    "total_rows": total_rows,
                    "total_columns": total_columns,
                    "preview_rows": preview_rows,
                    "truncated": sheet_truncated,
                }
            )
            workbook_truncated = workbook_truncated or sheet_truncated

        return {
            "filename": filename,
            "selected_sheet_name": selected_worksheet.title,
            "file_size_bytes": len(content),
            "sheet_count": sheet_count,
            "previewed_sheet_count": len(sheets),
            "sheets": sheets,
            "truncated": workbook_truncated,
        }
    except HTTPException:
        raise
    except (
        BadZipFile,
        EOFError,
        InvalidFileException,
        KeyError,
        OSError,
        SyntaxError,
        TypeError,
        ValueError,
    ) as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="上传的 .xlsx 工作簿损坏或无法读取",
        ) from exc
    finally:
        if workbook is not None:
            workbook.close()


def _import_cell_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, bool):
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else str(value)
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    text = str(value)
    if len(text) > MAX_IMPORT_CELL_CHARACTERS:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="Excel 单元格内容过长，无法安全准备导入",
        )
    return text


def _unique_headers(values: tuple[Any, ...]) -> list[str]:
    counts: dict[str, int] = {}
    headers: list[str] = []
    for index, value in enumerate(values):
        base = str(value).strip() if value is not None else ""
        base = base or f"列{index + 1}"
        if len(base) > 200:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="Excel 表头过长，无法安全准备导入",
            )
        counts[base] = counts.get(base, 0) + 1
        headers.append(base if counts[base] == 1 else f"{base}_{counts[base]}")
    return headers


def _parse_import_content(
    content: bytes,
    filename: str,
    sheet_name: str | None = None,
) -> dict:
    workbook = None
    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        worksheet = _selected_worksheet(workbook, sheet_name)
        if worksheet.max_row is None or worksheet.max_column is None:
            worksheet.calculate_dimension(force=True)
        total_rows = int(worksheet.max_row or 0)
        total_columns = int(worksheet.max_column or 0)
        if total_columns < 1 or total_rows < 1:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="选中的 Excel 工作表为空",
            )
        if total_columns > MAX_IMPORT_COLUMNS:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=f"Excel 导入最多支持 {MAX_IMPORT_COLUMNS} 列",
            )
        if total_rows - 1 > MAX_IMPORT_ROWS:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=f"Excel 导入最多支持 {MAX_IMPORT_ROWS} 条数据，请拆分文件后重试",
            )

        values = list(
            worksheet.iter_rows(
                min_row=1,
                max_row=total_rows,
                min_col=1,
                max_col=total_columns,
                values_only=True,
            )
        )
        headers = _unique_headers(values[0])
        rows: list[list[Any]] = []
        text_characters = sum(len(header) for header in headers)
        for source_row in values[1:]:
            if all(value is None or str(value).strip() == "" for value in source_row):
                continue
            row = [_import_cell_value(value) for value in source_row]
            text_characters += sum(len(str(value)) for value in row)
            if text_characters > MAX_IMPORT_TOTAL_TEXT_CHARACTERS:
                raise HTTPException(
                    status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                    detail="Excel 可导入文本总量超过安全限制，请拆分文件后重试",
                )
            rows.append(row)

        return {
            "filename": filename,
            "sheet_name": worksheet.title,
            "sheet_count": len(workbook.worksheets),
            "headers": headers,
            "rows": rows,
        }
    except HTTPException:
        raise
    except (
        BadZipFile,
        EOFError,
        InvalidFileException,
        KeyError,
        OSError,
        SyntaxError,
        TypeError,
        ValueError,
    ) as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="上传的 .xlsx 工作簿损坏或无法读取",
        ) from exc
    finally:
        if workbook is not None:
            workbook.close()


def _parse_excel_bundle(
    content: bytes,
    filename: str,
    sheet_name: str | None,
) -> dict:
    return {
        "preview": _parse_excel_content(content, filename, sheet_name),
        "import_data": _parse_import_content(content, filename, sheet_name),
    }


async def _read_excel_upload(file: UploadFile) -> tuple[bytes, str]:
    filename = _safe_filename(file.filename)
    if not filename.lower().endswith(".xlsx"):
        await file.close()
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail="目前仅支持 .xlsx 格式的 Excel 文件",
        )

    try:
        if file.size is not None and file.size > MAX_FILE_BYTES:
            raise HTTPException(
                status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                detail="Excel 文件不能超过 5 MiB",
            )
        content = await file.read(MAX_FILE_BYTES + 1)
    finally:
        await file.close()

    if len(content) > MAX_FILE_BYTES:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail="Excel 文件不能超过 5 MiB",
        )

    return content, filename


async def parse_excel_upload(file: UploadFile) -> dict:
    """Read a bounded .xlsx preview without persisting the uploaded file."""
    content, filename = await _read_excel_upload(file)

    return await run_in_threadpool(_parse_excel_content, content, filename)


async def parse_excel_upload_bundle(
    file: UploadFile,
    sheet_name: str | None = None,
) -> dict:
    """Read one upload into both a chat preview and a bounded import dataset."""
    content, filename = await _read_excel_upload(file)
    return await run_in_threadpool(_parse_excel_bundle, content, filename, sheet_name)
